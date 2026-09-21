"""Independent acceptance coverage for Phase 15 import and export APIs."""

from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework_simplejwt.tokens import AccessToken

from academics.models import Course
from accounts.models import User, UserRole
from scheduling.models import Schedule, ScheduleEntry, ScheduleVersion
from tests.test_phase7_validator_acceptance import ready_graph as phase7_ready_graph


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF = "application/pdf"
TEMPLATE_URL = "/api/imports/semester-plan/template/"
VALIDATE_URL = "/api/imports/semester-plan/validate/"
APPLY_URL = "/api/imports/semester-plan/apply/"
SHEETS = {
    "README", "Courses", "StudentGroups", "CourseOfferings", "TeachingComponents",
    "ComponentGroups", "TeachingAssignments", "RoomRequirements",
    "RequirementCapabilities",
}
EXPORT_SHEETS = {
    "Timetable", "Analytics Summary", "Department Load", "Instructor Workload",
    "Room Utilization", "Student Group Load", "Quality",
}


@pytest.fixture
def ready_graph(db):
    return phase7_ready_graph.__wrapped__(db)


def authenticate(client, user):
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {AccessToken.for_user(user)}")


def user(role, department=None):
    return User.objects.create_user(
        username=f"phase15-{role.lower()}-{User.objects.count()}",
        password="secret-pass-123",
        role=role,
        department=department,
    )


def publish_college_version(client, semester):
    assert client.post(
        "/api/schedules/generate-college-draft/", {"semester": semester.pk}, format="json"
    ).status_code == 200
    version = ScheduleVersion.objects.get(version_number=1)
    for action in ("submit", "review", "approve", "publish"):
        assert client.post(f"/api/schedule-versions/{version.pk}/{action}/", {}, format="json").status_code == 200
    return version


def template_bytes(client):
    response = client.get(TEMPLATE_URL)
    assert response.status_code == 200, response.content
    assert response["Content-Type"].split(";")[0] == XLSX
    assert "attachment" in response["Content-Disposition"]
    return response.content


def workbook_with_course(template, *, formula=False):
    workbook = load_workbook(BytesIO(template))
    course = workbook["Courses"]
    course.append(["C01", "P15-COURSE", "=SUM(1,2)" if formula else "Phase 15 Course"])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def upload(client, url, graph, content, name="plan.xlsx"):
    return client.post(
        url,
        {
            "department": str(graph["departments"]["A"].pk),
            "semester": str(graph["semester"].pk),
            "file": SimpleUploadedFile(name, content, content_type=XLSX),
        },
        format="multipart",
    )


def test_exports_are_valid_scoped_documents_and_do_not_write(api_client, ready_graph):
    admin = user(UserRole.COLLEGE_ADMIN)
    authenticate(api_client, admin)
    version = publish_college_version(api_client, ready_graph["semester"])
    version.refresh_from_db()
    before = (
        Schedule.objects.count(), ScheduleVersion.objects.count(), ScheduleEntry.objects.count(),
        version.status, version.created_at,
    )
    routes = (
        f"/api/schedule-versions/{version.pk}/export/xlsx/",
        "/api/published-schedules/current/export/xlsx/?semester=" + str(ready_graph["semester"].pk),
    )
    for route in routes:
        response = api_client.get(route)
        assert response.status_code == 200, response.content
        assert response["Content-Type"].split(";")[0] == XLSX
        assert "attachment" in response["Content-Disposition"]
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        assert set(workbook.sheetnames) == EXPORT_SHEETS
        assert workbook["Timetable"].max_row == 2
        assert all(cell.data_type != "f" for row in workbook["Timetable"].iter_rows() for cell in row)
    for route in (
        f"/api/schedule-versions/{version.pk}/export/pdf/",
        "/api/published-schedules/current/export/pdf/?semester=" + str(ready_graph["semester"].pk),
    ):
        response = api_client.get(route)
        assert response.status_code == 200, response.content
        assert response["Content-Type"].split(";")[0] == PDF
        assert response.content.startswith(b"%PDF")
    version.refresh_from_db()
    assert (
        Schedule.objects.count(), ScheduleVersion.objects.count(), ScheduleEntry.objects.count(),
        version.status, version.created_at,
    ) == before


def test_export_formula_values_are_neutralized(api_client, ready_graph):
    ready_graph["component"].offering.course.name = "=SUM(1,2)"
    ready_graph["component"].offering.course.save(update_fields=["name"])
    admin = user(UserRole.COLLEGE_ADMIN)
    authenticate(api_client, admin)
    version = publish_college_version(api_client, ready_graph["semester"])
    response = api_client.get(f"/api/schedule-versions/{version.pk}/export/xlsx/")
    assert response.status_code == 200
    values = [cell.value for row in load_workbook(BytesIO(response.content), data_only=False)["Timetable"].iter_rows() for cell in row]
    assert "'=SUM(1,2)" in values


def test_published_export_uses_authoritative_publication_and_handles_arabic(api_client, ready_graph):
    admin = user(UserRole.COLLEGE_ADMIN)
    authenticate(api_client, admin)
    missing = api_client.get(
        "/api/published-schedules/current/export/xlsx/",
        {"semester": ready_graph["semester"].pk},
    )
    assert missing.status_code == 404
    ready_graph["component"].offering.course.name = "مقدمة في الرياضيات"
    ready_graph["component"].offering.course.save(update_fields=["name"])
    version = publish_college_version(api_client, ready_graph["semester"])
    response = api_client.get(f"/api/schedule-versions/{version.pk}/export/pdf/")
    assert response.status_code == 200, response.content
    assert response.content.startswith(b"%PDF")


def test_template_validation_is_read_only_and_apply_is_create_only(api_client, ready_graph):
    admin = user(UserRole.COLLEGE_ADMIN)
    authenticate(api_client, admin)
    template = template_bytes(api_client)
    assert set(load_workbook(BytesIO(template)).sheetnames) == SHEETS
    workbook = workbook_with_course(template)
    before = Course.objects.count()
    first = upload(api_client, VALIDATE_URL, ready_graph, workbook)
    second = upload(api_client, VALIDATE_URL, ready_graph, workbook)
    assert first.status_code == second.status_code == 200
    assert first.json()["valid"] is second.json()["valid"] is True
    assert Course.objects.count() == before
    applied = upload(api_client, APPLY_URL, ready_graph, workbook)
    assert applied.status_code == 200, applied.content
    assert applied.json()["applied"] is True
    assert applied.json()["created"]["courses"] == 1
    assert Course.objects.count() == before + 1
    repeat = upload(api_client, APPLY_URL, ready_graph, workbook)
    assert repeat.status_code == 400
    assert repeat.json()["applied"] is False
    assert Course.objects.count() == before + 1


def test_import_refuses_formula_and_foreign_department(api_client, ready_graph):
    own = ready_graph["departments"]["A"]
    department_admin = user(UserRole.DEPARTMENT_ADMIN, own)
    authenticate(api_client, department_admin)
    template = template_bytes(api_client)
    formula = upload(api_client, VALIDATE_URL, ready_graph, workbook_with_course(template, formula=True))
    assert formula.status_code == 200
    assert formula.json()["valid"] is False
    assert any(issue["code"] == "FORMULA_NOT_ALLOWED" for issue in formula.json()["issues"])
    foreign = api_client.post(
        VALIDATE_URL,
        {"department": str(ready_graph["departments"]["B"].pk), "semester": str(ready_graph["semester"].pk), "file": SimpleUploadedFile("plan.xlsx", template, content_type=XLSX)},
        format="multipart",
    )
    assert foreign.status_code == 400
    for role in (UserRole.SCHEDULER, UserRole.VIEWER, UserRole.INSTRUCTOR):
        authenticate(api_client, user(role, own))
        assert api_client.get(TEMPLATE_URL).status_code == 403
        assert upload(api_client, VALIDATE_URL, ready_graph, template).status_code == 403
